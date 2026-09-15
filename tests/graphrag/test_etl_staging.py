from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.graphrag.etl_staging import read_table_rows
from app.graphrag.schema_etl_row_processing import RowProcessingError


def _write_xls(path: Path, rows: list[list[object]]) -> None:
    """用 xlwt 造一个最小的 xls（旧版二进制 Excel）夹具——xlrd 只能读、
    xlwt 只能写，生产代码不导入 xlwt。"""
    import xlwt

    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Sheet1")
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            worksheet.write(row_idx, col_idx, value)
    workbook.save(str(path))


def test_read_table_rows_reads_csv_with_utf8_bom(tmp_path: Path):
    """Excel 的 CSV UTF-8 导出会在文件开头写 BOM。utf-8-sig 会剥掉它；
    如果退化成纯 utf-8，BOM 会粘在第一个列名前面，整份文件被误判成缺列。"""
    path = tmp_path / "with_bom.csv"
    path.write_bytes("名称,编号\n抹茶,A1\n".encode("utf-8-sig"))

    rows = list(read_table_rows(path))

    assert rows == [{"名称": "抹茶", "编号": "A1"}]


def test_read_table_rows_falls_back_to_gbk(tmp_path: Path):
    """国内 Excel 导出的 CSV 常见默认编码是 GBK，UTF-8 严格解码会失败。"""
    path = tmp_path / "gbk.csv"
    path.write_bytes("名称,编号\n抹茶,A1\n".encode("gbk"))

    rows = list(read_table_rows(path))

    assert rows == [{"名称": "抹茶", "编号": "A1"}]


def test_read_table_rows_reads_tsv_with_tab_delimiter(tmp_path: Path):
    path = tmp_path / "data.tsv"
    path.write_text("name\tcode\nfoo\tA1\n", encoding="utf-8")

    rows = list(read_table_rows(path))

    assert rows == [{"name": "foo", "code": "A1"}]


def test_read_table_rows_skips_xlsx_phantom_all_empty_rows(tmp_path: Path):
    """openpyxl 的 read_only 迭代会按已用范围补齐行数，被清空的行会
    产出全空的幽灵行。这些行必须安静跳过，不能当成缺列的脏数据。"""
    path = tmp_path / "phantom.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "code"])
    sheet.append(["foo", "A1"])
    sheet.append(["bar", "A2"])
    sheet["A3"] = None
    sheet["B3"] = None
    workbook.save(path)

    rows = list(read_table_rows(path))

    assert rows == [{"name": "foo", "code": "A1"}]


def test_read_table_rows_on_empty_xlsx_sheet_yields_nothing(tmp_path: Path):
    path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.active.append(["name", "code"])
    workbook.save(path)

    assert list(read_table_rows(path)) == []


def test_read_table_rows_rejects_unsupported_suffix(tmp_path: Path):
    path = tmp_path / "data.json"
    path.write_text("{}", encoding="utf-8")

    with pytest.raises(RowProcessingError, match="不支持的数据文件类型"):
        list(read_table_rows(path))


def test_read_table_rows_is_lazy(tmp_path: Path):
    """staging 必须保持流式：真实规模是 MUJI 一张 SKU 表 18 万+ 行。
    拿到迭代器时文件还没被遍历完，只有真正迭代才逐行产出。

    注意：这条断言只证明"迭代器能产出正确的第一行"，不能单独证明它是
    真流式——一个先把整份文件读完再包成 iter() 的急切实现，同样能让
    这条断言通过。真正能区分二者的判别在
    test_read_table_rows_does_not_touch_file_until_iterated 里，两条
    测试互补，缺一不可。"""
    path = tmp_path / "data.csv"
    path.write_text("name\nfoo\nbar\n", encoding="utf-8")

    iterator = read_table_rows(path)
    first = next(iterator)

    assert first == {"name": "foo"}


def test_read_table_rows_does_not_touch_file_until_iterated(tmp_path: Path):
    """真正的生成器函数在被调用时只是创建一个生成器对象，函数体一行都
    不会执行；只有第一次 next() 才会真正跑到第一条语句。用一个不存在的
    路径来验证这一点：如果 read_table_rows 是真生成器，调用它本身不应该
    抛 FileNotFoundError，只有 next() 才应该抛；如果它退化成"先把整个
    文件读完（或先打开文件）再包成迭代器返回"（比如 return iter(list(...))
    这种急切实现），文件不存在会在调用点本身就抛出来。

    这个退化正是 18 万+ 行场景下最需要防住的回归——一旦变成急切求值，
    "流式、不整份文件常驻内存"的约束就名存实亡了，而只断言"能拿到第一
    行"（见 test_read_table_rows_is_lazy）抓不住这种退化，因为急切实现
    同样能给出正确的第一行。"""
    missing_path = tmp_path / "does_not_exist.csv"

    iterator = read_table_rows(missing_path)  # 调用本身不应该抛异常

    with pytest.raises(FileNotFoundError):
        next(iterator)


# ── 重名列 ────────────────────────────────────────────────────────────
#
# 真实事故形态：把表头直接当 dict 键，后出现的同名列会静默覆盖先出现的，
# 整列数据消失，而跑批报告仍然是"成功"。真实的表里重名列很常见——MUJI 的
# SKU 主数据表有 Color×2（英文/现地语）、Size×2、Retail price×4、
# Depth/Width/Height/Weight 各 ×3（ボール/ケース/ピース）。


def test_read_table_rows_keeps_both_columns_when_csv_header_repeats(tmp_path: Path):
    path = tmp_path / "dup.csv"
    path.write_text("Color,Color\nNatural,ナチュラル\n", encoding="utf-8")

    rows = list(read_table_rows(path))

    assert rows == [{"Color": "Natural", "Color (2)": "ナチュラル"}]


def test_read_table_rows_keeps_both_columns_when_xlsx_header_repeats(tmp_path: Path):
    path = tmp_path / "dup.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Color", "Color"])
    sheet.append(["Natural", "ナチュラル"])
    workbook.save(path)

    rows = list(read_table_rows(path))

    assert rows == [{"Color": "Natural", "Color (2)": "ナチュラル"}]


def test_read_table_rows_keeps_both_columns_when_xls_header_repeats(tmp_path: Path):
    path = tmp_path / "dup.xls"
    _write_xls(path, [["Color", "Color"], ["Natural", "ナチュラル"]])

    rows = list(read_table_rows(path))

    assert rows == [{"Color": "Natural", "Color (2)": "ナチュラル"}]


def test_read_table_rows_numbers_three_way_duplicates_in_order(tmp_path: Path):
    """三列同名时后缀按出现顺序递增，第一列保持原名——用户在字段映射里
    看到的 "Depth"、"Depth (2)"、"Depth (3)" 才能跟表里从左到右的
    ボール/ケース/ピース 对上。"""
    path = tmp_path / "triple.csv"
    path.write_text("Depth,Depth,Depth\n10,20,30\n", encoding="utf-8")

    rows = list(read_table_rows(path))

    assert rows == [{"Depth": "10", "Depth (2)": "20", "Depth (3)": "30"}]


def test_read_table_rows_leaves_unique_column_names_untouched(tmp_path: Path):
    """没有重名时一个后缀都不能加：列名变了，存着的字段映射就全部对不上，
    用户会看到一份"配置丢了"的表。"""
    path = tmp_path / "unique.csv"
    path.write_text("Color,Size\nNatural,S\n", encoding="utf-8")

    rows = list(read_table_rows(path))

    assert rows == [{"Color": "Natural", "Size": "S"}]


def test_read_table_rows_avoids_colliding_with_an_existing_suffixed_name(tmp_path: Path):
    """表头里本来就有一列叫 "Color (2)" 时，算出来的后缀不能直接用——
    否则第三列又把第二列覆盖掉，回到这个 bug 本身。"""
    path = tmp_path / "collide.csv"
    path.write_text("Color,Color (2),Color\na,b,c\n", encoding="utf-8")

    rows = list(read_table_rows(path))

    assert rows == [{"Color": "a", "Color (2)": "b", "Color (3)": "c"}]


def test_read_table_rows_keeps_unnamed_columns_apart(tmp_path: Path):
    """空列名彼此也是重名。MUJI 那张表有几十列表头是空的，全部挤进同一个
    "" 键的话，它们会互相覆盖成一列。"""
    path = tmp_path / "unnamed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", None, None])
    sheet.append(["foo", "x", "y"])
    workbook.save(path)

    rows = list(read_table_rows(path))

    assert rows == [{"name": "foo", "": "x", " (2)": "y"}]


# CSV 路径为了做重名列去重，从 csv.DictReader 换成了 csv.reader + 手工建
# 字典。DictReader 有两个不写在签名里的行为，换掉之后必须原样保住，否则
# 会从"丢一列"退化成"丢一整行"或"整份文件报缺列"。


def test_read_table_rows_skips_blank_csv_lines(tmp_path: Path):
    path = tmp_path / "blank.csv"
    path.write_text("name,code\nfoo,A1\n\nbar,A2\n", encoding="utf-8")

    rows = list(read_table_rows(path))

    assert rows == [{"name": "foo", "code": "A1"}, {"name": "bar", "code": "A2"}]


def test_read_table_rows_pads_short_csv_rows_with_none(tmp_path: Path):
    """值比表头少的行补 None（DictReader 的 restval 默认值）。补成 "" 的话，
    下游分不清"这一格是空的"和"这一行根本没有这一列"。"""
    path = tmp_path / "short.csv"
    path.write_text("name,code,note\nfoo,A1\n", encoding="utf-8")

    rows = list(read_table_rows(path))

    assert rows == [{"name": "foo", "code": "A1", "note": None}]
